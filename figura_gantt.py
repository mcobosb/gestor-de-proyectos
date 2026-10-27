"""Genera una figura de diagrama de Gantt (PNG + PDF) a partir del Excel de
planificación creado por bc3_to_gantt.py.

Uso:
    python gantt_figura_presentacion.py [planificacion.xlsx] [salida_sin_extension]

Sin archivo de entrada, busca un único .xlsx en la carpeta del script.
Si no hay ninguno o hay varios, muestra un error y no genera las figuras.
Se cuentan todos los .xlsx, incluidas las bases de precios; las copias .bak
y los archivos temporales con otra extensión no se cuentan.
Los nombres de salida se deducen del Excel seleccionado.

El Excel contiene fórmulas; para leer las fechas calculadas se recalcula una
copia temporal con LibreOffice (o se usan los valores guardados por Excel).

Versión: 1.0.0
Autor: Manuel Cobos Budia, Ingeniero de Caminos, Canales y Puertos Municipal,
Servicio de Infraestructuras y Obras del Excelentísimo Ayuntamiento de Lucena.
"""

from __future__ import annotations

import re
import shutil
import subprocess
import sys
import tempfile
import textwrap
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Dict, List, Optional

import matplotlib

matplotlib.use("Agg")
import matplotlib.dates as mdates
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, Patch, Polygon, Rectangle
from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent
GG_RATE = 0.13
BI_RATE = 0.06
IVA_RATE = 0.21  # Tipo general; ajustar si al contrato le corresponde otro tipo.
SOFFICE_CANDIDATES = [
    Path(r"C:\Program Files\LibreOffice\program\soffice.exe"),
    Path(r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"),
]

PALETTE = ["#1f77b4", "#e4572e", "#29a19c", "#f3a712", "#7b5ea7",
           "#2e86ab", "#c03221", "#5c946e", "#a23b72", "#3d5a80"]
DARK = "#1b2a41"
GRID = "#d9dee5"
WEEKEND = "#f2f4f7"
HOLIDAY = "#fde2e1"
FONT = "Segoe UI"
MONTHS_ES = ["ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sep", "oct", "nov", "dic"]


def find_unique_input(directory: Path, suffix: str) -> Path:
    """Localiza un único fichero de la extensión indicada o explica la ambigüedad."""
    candidates = sorted(
        (
            path for path in directory.iterdir()
            if path.is_file() and path.suffix.lower() == suffix.lower()
        ),
        key=lambda path: path.name.lower(),
    )
    if not candidates:
        raise SystemExit(f"Error: no se encontró ningún archivo {suffix} en {directory}.")
    if len(candidates) > 1:
        names = "\n".join(f"  - {path.name}" for path in candidates)
        raise SystemExit(
            f"Error: se encontró más de un archivo {suffix} en {directory}:\n{names}\n"
            "Deja un único archivo de entrada o indica su ruta como argumento."
        )
    return candidates[0]


@dataclass
class Activity:
    id: int
    chapter: str
    code: str
    summary: str
    unit: str
    quantity: Optional[float]
    amount: float
    duration: int
    start: date
    end: date
    continuous: bool


def to_date(v) -> Optional[date]:
    """Convierte fechas, instantes y números de serie de Excel a una fecha de calendario."""
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)) and v > 20000:  # número de serie de Excel
        return date(1899, 12, 30) + timedelta(days=int(v))
    return None


def recalculated_copy(xlsx: Path, tmpdir: Path) -> Path:
    """Solicita a LibreOffice recalcular una copia del libro si el ejecutable está disponible."""
    soffice = next((p for p in SOFFICE_CANDIDATES if p.exists()), None) or shutil.which("soffice")
    if not soffice:
        return xlsx
    subprocess.run([str(soffice), "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(tmpdir),
                    str(xlsx)], check=True, capture_output=True, timeout=180)
    out = tmpdir / xlsx.name
    return out if out.exists() else xlsx


def read_plan(xlsx: Path):
    """Lee fechas, actividades y parámetros calculados del Excel de planificación."""
    with tempfile.TemporaryDirectory() as tmp:
        src = recalculated_copy(xlsx, Path(tmp))
        wb = load_workbook(src, data_only=True)
        cfg = wb["Config"]
        origin = str(cfg["A2"].value or "")
        source_match = re.search(r"Origen:\s*(.+?)\s+.\s+generado el", origin)
        if not source_match:
            wb.close()
            raise SystemExit(f"No se pudo identificar el BC3 de origen en {xlsx}.")
        bc3_path = Path(source_match.group(1))
        if not bc3_path.is_absolute():
            bc3_path = xlsx.parent / bc3_path
        plan = str(cfg["B12"].value or "Completo").strip().capitalize()
        if plan not in {"Resumen", "Completo"}:
            wb.close()
            raise SystemExit(f"Plan de figura Gantt no válido en Config!B12: {plan!r}. Usa Resumen o Completo.")
        start_works, end_works = to_date(cfg["B4"].value), to_date(cfg["B9"].value)
        plazo, pem = cfg["B8"].value, cfg["B11"].value
        holidays = {to_date(cfg[f"A{r}"].value) for r in range(15, cfg.max_row + 1)} - {None}

        names: Dict[str, str] = {}
        ws = wb["Capítulos"]
        for r in range(2, ws.max_row + 1):
            if ws[f"A{r}"].value:
                names[str(ws[f"A{r}"].value)] = str(ws[f"B{r}"].value or "")

        acts: List[Activity] = []
        ws = wb["Planificación"]
        for r in range(6, ws.max_row + 1):
            aid, s, e = ws[f"A{r}"].value, to_date(ws[f"AA{r}"].value), to_date(ws[f"AB{r}"].value)
            if aid in (None, "") or s is None or e is None:
                continue
            acts.append(Activity(int(aid), str(ws[f"B{r}"].value or "Otros"), str(ws[f"C{r}"].value or ""),
                                 str(ws[f"D{r}"].value or ""), str(ws[f"E{r}"].value or ""),
                                 ws[f"F{r}"].value if isinstance(ws[f"F{r}"].value, (int, float)) else None,
                                 float(ws[f"G{r}"].value or 0), int(ws[f"AC{r}"].value or 0), s, e,
                                 str(ws[f"L{r}"].value or "").upper() == "S"))
        wb.close()
        title = read_bc3_title(bc3_path)
    if not acts:
        raise SystemExit("No hay fechas calculadas en el Excel. Ábrelo y guárdalo en Excel, o instala LibreOffice.")
    if start_works is None:
        start_works = min(a.start for a in acts)
    if end_works is None:
        end_works = max(a.end for a in acts)
    return acts, names, start_works, end_works, plazo, pem, holidays, title, plan


def read_bc3_title(path: Path) -> str:
    """Lee la codificación FIEBDC y extrae el título del concepto raíz del presupuesto."""
    if not path.is_file():
        raise SystemExit(f"No se encuentra el BC3 de origen: {path}")
    raw = path.read_bytes()
    head = raw[:300].decode("latin-1", errors="ignore").upper()
    encoding = "cp850" if "|850|" in head else "cp437" if "|437|" in head else "cp1252"
    text = raw.decode(encoding, errors="replace")
    for record in text.split("~"):
        fields = record.rstrip("\r\n").split("|")
        if (fields[0].strip().upper() == "C" and len(fields) > 3
                and fields[1].split("\\")[0].strip().endswith("##")):
            title = fields[3].strip()
            if title:
                return title
    raise SystemExit(f"No se encontró el título del concepto raíz en el BC3: {path}")


def fmt_date(d: date) -> str:
    """Formatea una fecha como día/mes/año con dos cifras en día y mes."""
    return f"{d.day:02d}/{d.month:02d}/{d.year}"


def fmt_eur(v: float) -> str:
    """Formatea un importe en euros con separadores españoles."""
    return f"{v:,.2f} €".replace(",", "X").replace(".", ",").replace("X", ".")


def short(text: str, n: int) -> str:
    """Normaliza los espacios y abrevia el texto a la longitud máxima indicada."""
    text = " ".join(text.split())
    return text if len(text) <= n else text[: n - 1] + "…"


def draw(acts, names, start_works, end_works, plazo, pem, holidays, title: str, plan: str, out_base: Path):
    """Dibuja el Gantt por capítulos o partidas y guarda las versiones PNG y PDF."""
    plt.rcParams.update({"font.family": FONT, "font.size": 8.5})

    chapters: List[str] = list(dict.fromkeys(a.chapter for a in acts))
    color = {ch: PALETTE[i % len(PALETTE)] for i, ch in enumerate(chapters)}

    rows = []  # (tipo, objeto)
    for ch in chapters:
        rows.append(("cap", ch))
        if plan == "Completo":
            rows.extend(("act", a) for a in acts if a.chapter == ch)
    n = len(rows)

    t0 = start_works - timedelta(days=start_works.weekday())
    t1 = max(max(a.end for a in acts), end_works) + timedelta(days=7)
    span = (t1 - t0).days

    fig_w = 17
    fig_h = max(fig_w * 9 / 16, 2.55 + 0.22 * n)
    fig = plt.figure(figsize=(fig_w, fig_h), dpi=200)
    fig.patch.set_facecolor("white")
    top_frac = 1.25 / fig_h
    bottom = 0.75 / fig_h
    ax = fig.add_axes([0.335, bottom, 0.645, 1 - top_frac - bottom - 0.55 / fig_h])
    ax.set_xlim(mdates.date2num(t0), mdates.date2num(t1))
    ax.set_ylim(n - 0.5, -0.5)

    # Fondo: fines de semana / festivos
    d = t0
    while d < t1:
        x = mdates.date2num(d)
        if d in holidays:
            ax.axvspan(x, x + 1, color=HOLIDAY, lw=0, zorder=0)
        elif d.weekday() >= 5:
            ax.axvspan(x, x + 1, color=WEEKEND, lw=0, zorder=0)
        d += timedelta(days=1)

    # Filas: bandas de capítulo
    for i, (kind, obj) in enumerate(rows):
        if kind == "cap":
            ax.axhspan(i - 0.5, i + 0.5, color=color[obj], alpha=0.10, lw=0, zorder=0.5)
        ax.axhline(i + 0.5, color=GRID, lw=0.4, zorder=0.6)

    # Barras
    bar_h = 0.56
    for i, (kind, obj) in enumerate(rows):
        if kind == "cap":
            ch_acts = [a for a in acts if a.chapter == obj]
            s = min(a.start for a in ch_acts)
            e = max(a.end for a in ch_acts) + timedelta(days=1)
            xs, xe = mdates.date2num(s), mdates.date2num(e)
            c = color[obj]
            h = 0.22
            ax.add_patch(Rectangle((xs, i - h / 2), xe - xs, h, color=DARK, zorder=3))
            for xx, sgn in ((xs, 1), (xe, -1)):
                ax.add_patch(Polygon([[xx, i - h / 2], [xx, i + 0.36], [xx + sgn * span * 0.006, i - h / 2]],
                                     closed=True, color=DARK, zorder=3))
            ax.add_patch(Rectangle((xs, i - h / 2 + 0.04), xe - xs, h - 0.08, color=c, zorder=3.1))
            amount = sum(a.amount for a in ch_acts)
            txt = f"{fmt_date(s)} → {fmt_date(e - timedelta(days=1))}  ·  {fmt_eur(amount)}"
            room_right = mdates.date2num(t1) - xe
            if room_right > span * 0.30:
                ax.text(xe + span * 0.008, i, txt, va="center", ha="left", fontsize=7.5, color=DARK,
                        fontweight="bold", zorder=4)
            elif xs - mdates.date2num(t0) > span * 0.30:
                ax.text(xs - span * 0.008, i, txt, va="center", ha="right", fontsize=7.5, color=DARK,
                        fontweight="bold", zorder=4)
            else:
                ax.text((xs + xe) / 2, i - 0.30, txt, va="center", ha="center", fontsize=7.3, color=DARK,
                        fontweight="bold", zorder=7,
                        bbox=dict(boxstyle="round,pad=0.25", fc="white", ec=c, lw=0.8))
            continue

        a: Activity = obj
        c = color[a.chapter]
        xs = mdates.date2num(a.start)
        xe = mdates.date2num(a.end + timedelta(days=1))
        if a.duration == 0:
            ax.scatter([xs + 0.5], [i], marker="D", s=40, color=DARK, zorder=4)
            continue
        box = FancyBboxPatch((xs, i - bar_h / 2), xe - xs, bar_h,
                             boxstyle="round,pad=0,rounding_size=0.15", mutation_aspect=1 / max(span, 1) * 40,
                             facecolor=c, edgecolor="white", lw=0.6, alpha=0.55 if a.continuous else 0.95,
                             hatch="////" if a.continuous else None, zorder=3)
        ax.add_patch(box)
        label = f"{a.duration} días laborables · {fmt_eur(a.amount)}" if a.continuous else f"{a.duration} d"
        width_days = xe - xs
        if width_days >= span * 0.035:
            ax.text((xs + xe) / 2, i, label, ha="center", va="center", fontsize=7,
                    color=DARK if a.continuous else "white", fontweight="bold", zorder=5,
                    bbox=dict(boxstyle="round,pad=0.15", fc="white", ec="none", alpha=0.85)
                    if a.continuous else None)
        else:
            ax.text(xe + span * 0.004, i, label, ha="left", va="center", fontsize=7, color=c,
                    fontweight="bold", zorder=5)

    # Líneas de inicio y fin de obra
    for dte, txt, col_ in ((start_works, "INICIO", "#2a9d8f"), (end_works, "FIN", "#c03221")):
        x = mdates.date2num(dte) + (1 if txt == "FIN" else 0)
        ax.axvline(x, color=col_, lw=1.4, ls=(0, (5, 3)), zorder=6)
        ax.text(x, n - 0.5 + 0.25, f"{txt} {fmt_date(dte)}", color="white", fontsize=7.5, fontweight="bold",
                va="top", ha="center", zorder=8, clip_on=False,
                bbox=dict(boxstyle="round,pad=0.3", fc=col_, ec="none"))

    # Eje temporal: lunes de cada semana, meses y años en bandas separadas
    ax.xaxis.set_major_locator(mdates.WeekdayLocator(byweekday=mdates.MO, interval=1 if span < 220 else 2))
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%d"))
    ax.tick_params(axis="x", labelsize=7, colors="#4a5568", length=0, pad=3)
    ax.tick_params(axis="x", which="both", top=True, labeltop=True, bottom=False, labelbottom=False)
    ax.grid(axis="x", color=GRID, lw=0.6, zorder=0.7)
    ax.set_yticks([])
    for s in ax.spines.values():
        s.set_visible(False)

    month_ax = fig.add_axes([ax.get_position().x0, ax.get_position().y1 + 0.17 / fig_h,
                             ax.get_position().width, 0.2 / fig_h])
    month_ax.set_xlim(ax.get_xlim())
    month_ax.set_ylim(0, 1)
    month_ax.axis("off")
    m = date(t0.year, t0.month, 1)
    k = 0
    month_width_inches = ax.get_position().width * fig_w
    while m < t1:
        nxt = date(m.year + (m.month == 12), m.month % 12 + 1, 1)
        a0, a1 = max(m, t0), min(nxt, t1)
        x0, x1 = mdates.date2num(a0), mdates.date2num(a1)
        month_ax.add_patch(Rectangle((x0, 0), x1 - x0, 1, color=DARK if k % 2 == 0 else "#33475b", lw=0))
        if (x1 - x0) / span * month_width_inches >= 0.24:
            month_ax.text((x0 + x1) / 2, 0.5, MONTHS_ES[m.month - 1].upper(), ha="center",
                          va="center", color="white", fontsize=8, fontweight="bold")
        m, k = nxt, k + 1

    year_ax = fig.add_axes([ax.get_position().x0, ax.get_position().y1 + 0.39 / fig_h,
                            ax.get_position().width, 0.16 / fig_h])
    year_ax.set_xlim(ax.get_xlim())
    year_ax.set_ylim(0, 1)
    year_ax.axis("off")
    year_ax.text(mdates.date2num(t0), 0.5, str(t0.year), ha="left", va="center",
                 color=DARK, fontsize=8, fontweight="bold")
    for year in range(t0.year + 1, t1.year + 1):
        year_start = date(year, 1, 1)
        if year_start < t1:
            x = mdates.date2num(year_start)
            year_ax.axvline(x, color=GRID, lw=0.8)
            year_ax.text(x, 0.5, str(year), ha="left", va="center", color=DARK,
                         fontsize=8, fontweight="bold")

    # Columna de texto (capítulos y unidades de obra)
    lab = fig.add_axes([0.012, ax.get_position().y0, 0.32, ax.get_position().height])
    lab.set_xlim(0, 1)
    lab.set_ylim(ax.get_ylim())
    lab.axis("off")
    hy = -0.6
    lab.text(0.0, hy, "CAPÍTULO" if plan == "Resumen" else "ID  ·  CAPÍTULO / UNIDAD DE OBRA",
             fontsize=8, fontweight="bold", color=DARK, va="bottom")
    if plan == "Completo":
        lab.text(0.80, hy, "MEDICIÓN", fontsize=8, fontweight="bold", color=DARK, va="bottom", ha="right")
        lab.text(0.985, hy, "DURAC.", fontsize=8, fontweight="bold", color=DARK, va="bottom", ha="right")
    for i, (kind, obj) in enumerate(rows):
        if kind == "cap":
            c = color[obj]
            lab.add_patch(Rectangle((0, i - 0.5), 1, 1, color=c, alpha=0.10, lw=0))
            lab.add_patch(Rectangle((0, i - 0.5), 0.012, 1, color=c, lw=0))
            ch_acts = [a for a in acts if a.chapter == obj]
            days = max(a.end for a in ch_acts) - min(a.start for a in ch_acts)
            chapter_title = names.get(obj, obj)
            lab.text(0.025, i, short(f"{obj.replace('#', '').upper()}  {chapter_title}", 62), va="center",
                     fontsize=8.5, fontweight="bold", color=c)
            if plan == "Completo":
                lab.text(0.985, i, f"{days.days + 1} d nat.", va="center", ha="right", fontsize=7.5,
                         fontweight="bold", color=c)
        else:
            a = obj
            lab.add_patch(Rectangle((0, i - 0.5), 0.012, 1, color=color[a.chapter], alpha=0.35, lw=0))
            lab.text(0.03, i, f"{a.id:>2}", va="center", fontsize=7.5, color="#718096")
            lab.text(0.075, i, short(f"{a.code}  {a.summary.capitalize()}", 60), va="center", fontsize=7.6,
                     color="#2d3748")
            if a.quantity is not None:
                qty = f"{a.quantity:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
                lab.text(0.80, i, f"{qty} {a.unit}", va="center", ha="right", fontsize=7.2, color="#4a5568")
            lab.text(0.985, i, "toda obra" if a.continuous else f"{a.duration} d", va="center", ha="right",
                     fontsize=7.4, color="#2d3748", fontweight="bold")
        lab.axhline(i + 0.5, color=GRID, lw=0.4)

    # Cabecera con datos clave
    head = fig.add_axes([0, 1 - top_frac, 1, top_frac])
    head.set_xlim(0, 1)
    head.set_ylim(0, 1)
    head.axis("off")
    head.add_patch(Rectangle((0, 0), 1, 1, color=DARK))
    title_size = max(8, 15 * min(1, 52 / max(len(title), 1)))
    head.text(0.015, 0.62, title, color="white", fontsize=title_size, fontweight="bold", va="center")
    head.text(0.015, 0.26,
              f"Diagrama de Gantt ({plan.lower()}) · PBL = PEM + GG {GG_RATE:.0%} + BI {BI_RATE:.0%} + IVA {IVA_RATE:.0%}",
              color="#cbd5e0", fontsize=8.5, va="center")
    pbl = pem * (1 + GG_RATE + BI_RATE) * (1 + IVA_RATE) if isinstance(pem, (int, float)) else None
    natural_days = (end_works - start_works).days + 1
    duration_months = natural_days / (365.2425 / 12)
    cards = [("INICIO DE OBRA", fmt_date(start_works), "#2a9d8f"),
             ("FIN DE OBRA", fmt_date(end_works), "#e4572e"),
             ("PLAZO", f"{plazo} d lab. · {natural_days} nat. · ~{duration_months:.1f} meses".replace(".", ",")
              if plazo else "-",
              "#f3a712"),
             ("PEM", fmt_eur(pem) if isinstance(pem, (int, float)) else "-", "#2e86ab"),
             ("PBL (IVA INCL.)", fmt_eur(pbl) if pbl is not None else "-", "#7b5ea7")]
    x = 0.35
    widths = [0.095, 0.095, 0.19, 0.105, 0.13]
    for (t, v, c), w in zip(cards, widths):
        head.add_patch(FancyBboxPatch((x, 0.14), w, 0.72, boxstyle="round,pad=0,rounding_size=0.02",
                                      facecolor="#ffffff12", edgecolor=c, lw=1.6))
        head.add_patch(Rectangle((x, 0.14), 0.005, 0.72, color=c))
        head.text(x + 0.009, 0.66, t, color=c, fontsize=7, fontweight="bold", va="center")
        head.text(x + 0.009, 0.36, v, color="white", fontsize=9.5, fontweight="bold", va="center")
        x += w + 0.008

    # Leyenda
    handles = [Patch(facecolor=color[ch], label=short(names.get(ch, ch).capitalize(), 34)) for ch in chapters]
    handles += [Patch(facecolor="#9aa5b1", hatch="////", alpha=0.6, label="Actividad continua (toda la obra)"),
                Patch(facecolor=WEEKEND, edgecolor="#cbd5e0", label="Fin de semana"),
                Patch(facecolor=HOLIDAY, label="Festivo")]
    fig.legend(handles=handles, loc="lower center", ncol=min(7, len(handles)), frameon=False, fontsize=7.5,
               bbox_to_anchor=(0.5, 0.0), handlelength=1.4, columnspacing=1.4)
    fig.text(0.985, 0.012, f"Generado {date.today():%d/%m/%Y}", ha="right", fontsize=6.5, color="#a0aec0")

    out_base.parent.mkdir(parents=True, exist_ok=True)
    png, pdf = out_base.with_suffix(".png"), out_base.with_suffix(".pdf")
    fig.savefig(png, dpi=200, facecolor="white")
    fig.savefig(pdf, facecolor="white")
    plt.close(fig)
    return png, pdf


def main(argv: List[str]) -> int:
    """Resuelve los argumentos de consola, lee el Excel y genera las figuras del plan."""
    xlsx = Path(argv[1]).resolve() if len(argv) > 1 else find_unique_input(BASE_DIR, ".xlsx")
    out = (Path(argv[2]).resolve() if len(argv) > 2 else
           xlsx.with_name(xlsx.stem.replace("_planificacion", "") + "_gantt_presentacion"))
    if not xlsx.exists():
        raise SystemExit(f"No existe {xlsx}. Ejecuta antes bc3_to_gantt.py")
    data = read_plan(xlsx)
    png, pdf = draw(*data, out_base=out)
    print(f"Generado: {png}\nGenerado: {pdf}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
